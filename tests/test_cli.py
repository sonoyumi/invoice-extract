"""End-to-end runs of the command on the files in examples/."""

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from invoice_extract.cli import run

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"
TODAY = date(2026, 9, 29)
MY_VAT = "IT02345678904"


def test_full_run_on_examples(tmp_path, capsys):
    out = tmp_path / "report.xlsx"
    assert run([str(EXAMPLES), "-o", str(out), "--my-vat", MY_VAT, "--csv", str(tmp_path / "csv")], today=TODAY) == 0
    text = capsys.readouterr().out
    assert "Файлов: 8 · счетов: 6 · дублей: 1 · не счета (пропущены): 1" in text
    assert "Ошибок: 3 · предупреждений: 2" in text
    assert "выставлен 10%" in text and "Studio Grafico Weiss" in text

    wb = load_workbook(out)
    fatture = list(wb["Fatture"].iter_rows(min_row=2, values_only=True))
    assert len(fatture) == 6
    assert {r[0] for r in fatture} == {"Emessa", "Ricevuta"}
    signed = next(r for r in fatture if r[1] == "TD06")  # the .p7m one
    assert signed[12] == 200 and signed[14] == 1020  # withholding and net to pay
    credit = next(r for r in fatture if r[1] == "TD04")
    assert credit[11] < 0
    assert (tmp_path / "csv" / "fatture.csv").exists()


def test_strict_mode_fails_on_errors(tmp_path):
    assert run([str(EXAMPLES), "-o", str(tmp_path / "r.xlsx"), "--strict"], today=TODAY) == 1


def test_clean_file_passes_strict(tmp_path):
    clean = EXAMPLES / "IT04567890126_00017.xml.p7m"
    assert run([str(clean), "-o", str(tmp_path / "r.xlsx"), "--strict"], today=TODAY) == 0


def test_date_filter(tmp_path, capsys):
    args = [str(EXAMPLES), "-o", str(tmp_path / "r.xlsx"), "--from", "2026-03-10", "--to", "2026-03-20"]
    assert run(args, today=TODAY) == 0
    assert "счетов: 3" in capsys.readouterr().out  # NC-0012, 17, 3


def test_italian_messages(tmp_path, capsys):
    assert run([str(EXAMPLES), "-o", str(tmp_path / "r.xlsx"), "--lang", "it"], today=TODAY) == 0
    out = capsys.readouterr().out
    assert "Errori: 3" in out and "TOTALE" in out and "Report salvato" in out
    problems = load_workbook(tmp_path / "r.xlsx")["Problemi"]
    assert any("aliquota 0% senza codice Natura" in (row[3] or "") for row in problems.iter_rows(values_only=True))


def test_wrong_input_gives_exit_code_2(tmp_path, capsys):
    assert run(["/no/such/dir", "-o", str(tmp_path / "r.xlsx")]) == 2
    assert run([str(tmp_path), "-o", str(tmp_path / "r.xlsx")]) == 2  # empty folder
    assert run([str(EXAMPLES), "-o", str(tmp_path / "r.xlsx"), "--my-vat", "IT01234567890"]) == 2
    err = capsys.readouterr().err
    assert "не найдены" in err and "Не найдено ни одного файла" in err and "неверная партита IVA" in err
    assert not (tmp_path / "r.xlsx").exists()
