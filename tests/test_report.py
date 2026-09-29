import csv

from factory import Body, L, invoice_xml
from openpyxl import load_workbook

from invoice_extract.models import Issue
from invoice_extract.parser import parse_invoices
from invoice_extract.report import write_csv, write_excel


def invoices():
    evil = Body(number="1", lines=[L('=HYPERLINK("http://x","clic")', "1", "100.00")])
    credit = Body(doc_type="TD04", number="NC1", lines=[L("-reso merce", "1", "10.00")])
    return parse_invoices(invoice_xml(evil), "a.xml") + parse_invoices(invoice_xml(credit), "b.xml")


def test_excel_has_all_sheets_and_signed_amounts(tmp_path):
    issue = Issue.make("a.xml", "error", "no_lines")
    path = write_excel(tmp_path / "out" / "r.xlsx", invoices(), [issue], lang="it")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Fatture", "Righe", "Riepilogo IVA", "Controparti", "Scadenze", "Problemi"]
    rows = list(wb["Fatture"].iter_rows(min_row=2, values_only=True))
    totals = [r[11] for r in rows]
    assert totals == [122, -12.2]
    assert wb["Fatture"]["L2"].number_format == "#,##0.00"
    assert wb["Problemi"]["A2"].value == "ERRORE" and "righe" in wb["Problemi"]["D2"].value


def test_formula_from_invoice_stays_text(tmp_path):
    wb = load_workbook(write_excel(tmp_path / "r.xlsx", invoices(), []))
    cell = wb["Righe"]["E2"]
    assert cell.value.startswith("=HYPERLINK") and cell.data_type == "s"


def test_csv_is_safe_and_italian_formatted(tmp_path):
    folder = write_csv(tmp_path / "csv", invoices(), [])
    with (folder / "righe.csv").open(encoding="utf-8-sig") as f:
        rows = list(csv.reader(f, delimiter=";"))
    assert rows[1][4].startswith("'=HYPERLINK") and rows[2][4] == "'-reso merce"
    assert rows[1][7] == "100,00" and rows[2][7] == "-10,00"
    assert rows[1][1] == "10/03/2026"
    assert (folder / "fatture.csv").exists() and (folder / "problemi.csv").exists()
